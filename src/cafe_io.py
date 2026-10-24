"""JSON, CSV and Excel adapters; spreadsheet contents are data, never executable formulas."""
from __future__ import annotations
import csv
import json
import sys
from datetime import date, timedelta
from pathlib import Path

from cafe_core import CafeError, save_json, validate_context, validate_history

VENDOR = Path(__file__).resolve().parents[1] / "vendor"
if VENDOR.exists():
    sys.path.insert(0, str(VENDOR))

LABELS = {
    "shop_id": "模拟店铺代号", "as_of_date": "填表日期", "sales_date": "昨日销量日期", "target_date": "预测日期（明天）",
    "cups_sold": "昨日杯数", "weather": "明日天气（未知也可）", "is_holiday": "明日是否节假日",
    "nearby_event": "明日周围活动", "owner_comment": "店主意见（假设，不是事实）", "daily_capacity": "明日最大产能（杯）",
    "grams_per_cup": "每杯咖啡粉/豆克数", "opening_stock_kg": "明日开店前预计可用存量 kg（不含另列在途）",
    "due_inbound_kg": "明日开店前确认到货的在途 kg", "safety_stock_kg": "店主设定的安全库存 kg",
    "pack_kg": "每包 kg", "max_order_kg": "单次采购上限 kg", "storage_capacity_kg": "储存上限 kg",
    "supplier_delivery_date": "新采购可在该日开店前到货", "closed": "明日是否休业", "stockout": "昨日是否售罄",
    "data_source": "比赛数据来源（固定 synthetic）"}
BOOL_KEYS = {"is_holiday", "closed", "stockout"}
INT_KEYS = {"cups_sold", "daily_capacity"}
FLOAT_KEYS = {"grams_per_cup", "opening_stock_kg", "due_inbound_kg", "safety_stock_kg", "pack_kg", "max_order_kg", "storage_capacity_kg"}
HISTORY_FIELDS = ["date", "cups", "weather", "is_holiday", "event", "stockout", "data_source"]


def demo_context(as_of="2026-10-21"):
    day = date.fromisoformat(as_of)
    return {"shop_id": "SYNTHETIC_CAFE_01", "as_of_date": as_of, "sales_date": (day-timedelta(days=1)).isoformat(),
            "target_date": (day+timedelta(days=1)).isoformat(), "cups_sold": 100, "weather": "模拟：晴，约 30°C",
            "is_holiday": False, "nearby_event": "模拟：没有已确认活动", "owner_comment": "模拟：希望减少囤货；没有额外客流证据",
            "daily_capacity": 300, "grams_per_cup": 18, "opening_stock_kg": 0.5, "due_inbound_kg": 0,
            "safety_stock_kg": 0.2, "pack_kg": 0.5, "max_order_kg": 10, "storage_capacity_kg": 20,
            "supplier_delivery_date": (day+timedelta(days=1)).isoformat(), "closed": False, "stockout": False, "data_source": "synthetic"}


def parse_field(key, value):
    if key in BOOL_KEYS:
        if type(value) is bool:
            return value
        if str(value).strip().lower() in ("true", "是", "1"):
            return True
        if str(value).strip().lower() in ("false", "否", "0"):
            return False
        raise CafeError(f"{key}请填写 true/false 或 是/否。")
    if key in INT_KEYS | FLOAT_KEYS:
        try:
            if isinstance(value, bool) or str(value).strip() == "":
                raise ValueError
            number = float(value)
            if key in INT_KEYS:
                if not number.is_integer():
                    raise ValueError
                return int(number)
            return number
        except (ValueError, TypeError, OverflowError):
            raise CafeError(f"{key}不是有效数字。") from None
    return "" if value is None else str(value)


def excel_module():
    try:
        import openpyxl
        return openpyxl
    except ImportError:
        raise CafeError("Excel 支持缺少 openpyxl。运行 python -m pip install -r requirements.txt；JSON/CSV 和主程序无需此依赖。") from None


def literal(cell, value):
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"  # Prevent formula injection, even with '=' in owner/model text.


def style_workbook(wb):
    from openpyxl.styles import Alignment, Font, PatternFill
    for ws in wb:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="174A50")
            cell.font = Font(color="FFFFFF", bold=True)
        for col in ws.columns:
            letter = col[0].column_letter
            ws.column_dimensions[letter].width = min(65, max(18, max(len(str(c.value or "")) for c in col[:30]) * 1.3))
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)


def add_rows(ws, rows):
    for row_index, row in enumerate(rows, 1):
        for col_index, value in enumerate(row, 1):
            literal(ws.cell(row_index, col_index), value)


def write_template(path, context=None, history=None):
    c = context or demo_context()
    wb = excel_module().Workbook()
    ws = wb.active
    ws.title = "输入"
    add_rows(ws, [["字段（勿改）", "值（可编辑）", "说明"]] + [[k, v, LABELS.get(k, k)] for k, v in c.items() if k != "weekday"])
    ws = wb.create_sheet("历史")
    add_rows(ws, [HISTORY_FIELDS] + [[r[k] for k in HISTORY_FIELDS] for r in (history or [])])
    ws = wb.create_sheet("说明")
    notes = [["事项", "说明"], ["数据", "全部内容必须是团队自建模拟数据；不可上传真实顾客或公司记录。"],
             ["库存时间", "opening_stock_kg 是明日开店前预计存量，已经扣除今天预计消耗；不是此刻库存。"],
             ["在途", "due_inbound_kg 仅填写确认在明日开店前到货、且没有计入 opening_stock_kg 的数量。"],
             ["保质期", "过期、受损、不适用的库存不得计入可用存量；本版不做食品安全判断。"],
             ["范围", "单店、单类咖啡、单日预测；不支持不同配方混算或多日交期推算。"],
             ["AI", "销量由模型估计；公斤与补货数量使用精确的业务算术；Excel 不自行调用模型。"],
             ["历史", "每行一天，日期不重复且不晚于昨日；昨日一行必须与输入表一致。"],
             ["禁止公式", "导入单元格必须是值，不接受公式或宏；日期请用文本 YYYY-MM-DD。"],
             ["限制", "默认样例不代表真实预测能力。十轮分歧范围不是置信区间。"]]
    add_rows(ws, notes)
    style_workbook(wb)
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def read_input(path):
    path = Path(path)
    if path.stat().st_size > 5_000_000:
        raise CafeError("输入文件超过 5 MB。")
    if path.suffix.lower() == ".json":
        data = json.loads(path.read_text(encoding="utf-8-sig"))
        if not isinstance(data, dict) or set(data) != {"context", "history"}:
            raise CafeError("JSON 顶层需要 context 和 history。")
        c = validate_context(data["context"])
        return c, validate_history(data["history"], c)
    if path.suffix.lower() != ".xlsx":
        raise CafeError("输入文件只支持 .json 或无宏 .xlsx。")
    # Limit ZIP expansion before passing a user-controlled workbook to openpyxl.
    import zipfile
    with zipfile.ZipFile(path) as z:
        if len(z.infolist()) > 200 or sum(i.file_size for i in z.infolist()) > 20_000_000:
            raise CafeError("Excel 解压后过大。")
    wb = excel_module().load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if not {"输入", "历史"} <= set(wb.sheetnames):
            raise CafeError("Excel 缺少“输入”或“历史”表。")
        for ws in wb:
            if ws.max_row > 400 or ws.max_column > 15:
                raise CafeError("Excel 行列超出模板范围。")
            if any(cell.data_type == "f" for row in ws for cell in row):
                raise CafeError("Excel 含公式，请粘贴为数值后导入。")
        c = {}
        for row in wb["输入"].iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            key = str(row[0])
            if key in c:
                raise CafeError("输入字段重复。")
            c[key] = parse_field(key, row[1])
        c = validate_context(c)
        values = list(wb["历史"].values)
        if list(values[0]) != HISTORY_FIELDS:
            raise CafeError("历史表列名或顺序与模板不符。")
        rows = []
        for row in values[1:]:
            if not any(v is not None for v in row):
                continue
            r = dict(zip(HISTORY_FIELDS, row))
            r["cups"] = parse_field("cups_sold", r["cups"])
            for key in ("is_holiday", "stockout"):
                r[key] = parse_field(key, r[key])
            for key in ("weather", "event"):
                r[key] = "" if r[key] is None else str(r[key])
            rows.append(r)
        return c, validate_history(rows, c)
    finally:
        wb.close()


def export_result(path, result, snapshot=None):
    path = Path(path)
    if path.suffix.lower() == ".json":
        save_json(path, {"result": result, "audit": snapshot or {}})
        return
    if path.suffix.lower() != ".xlsx":
        raise CafeError("结果导出支持 .xlsx 或 .json。")
    wb = excel_module().Workbook()
    wb.active.title = "建议（待人工核对）"
    summary = [["指标", "值"], ["状态", result["status"]], ["运行编号", result["run_id"]],
               ["数据来源", "团队模拟数据"], ["模型", result["model"].get("model")],
               ["模型摘要", result["model"].get("digest")], ["提示词版本", result["model"].get("prompt_version")]]
    summary += [[k, v] for k, v in result.get("summary", {}).items() if not isinstance(v, list)]
    summary += [[k, v] for k, v in result.get("procurement", {}).items() if not isinstance(v, list)]
    summary += [["注意", w] for w in result.get("warnings", []) + result.get("procurement", {}).get("blockers", [])]
    add_rows(wb.active, summary)
    add_rows(wb.create_sheet("输入快照"), [["字段", "值"]] + list(result["context"].items()))
    rows = [["轮次", "随机种子", "估计杯数", "模型理由", "引用证据", "局限"]]
    rows += [[i, s["seed"], s["forecast_cups"], s["reason"], ", ".join(s["evidence_ids"]), "；".join(s["limitations"])] for i, s in enumerate(result.get("samples", []), 1)]
    add_rows(wb.create_sheet("十轮样本"), rows)
    add_rows(wb.create_sheet("模型参考证据"), [["证据编号", "传给模型的内容"]] +
             [[f["id"], json.dumps(f, ensure_ascii=False)] for f in result.get("evidence", [])])
    if snapshot:
        for name, key in (("审批订单", "orders"), ("操作记录", "events")):
            data = snapshot[key]
            if data:
                fields = list(data[0])
                add_rows(wb.create_sheet(name), [fields] + [[r[k] for k in fields] for r in data])
    style_workbook(wb)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def safe_csv(path, rows):
    if not rows:
        raise CafeError("没有可导出的行。")
    with Path(path).open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]))
        writer.writeheader()
        for row in rows:
            writer.writerow({k: ("'" + v if isinstance(v, str) and v.lstrip().startswith(("=", "+", "-", "@")) else v) for k, v in row.items()})
